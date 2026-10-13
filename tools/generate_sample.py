from pptx import Presentation
prs = Presentation()
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "ISHA Poster"
subtitle.text = "Sample presentation"
notes_slide = slide.notes_slide
notes_slide.notes_text_frame.text = "Good morning everyone. Today I'd like to present our study."

bullet_slide_layout = prs.slide_layouts[1]
slide2 = prs.slides.add_slide(bullet_slide_layout)
shapes = slide2.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Background"
tf = body_shape.text_frame
tf.text = "Relaxin has been reported to affect ligament laxity"
notes_slide2 = slide2.notes_slide
notes_slide2.notes_text_frame.text = "First, let me explain the background. Relaxin has been reported to affect RXFP1 expression in SOX9-positive chondrocytes."

prs.save("sample_poster.pptx")
